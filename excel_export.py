import pandas as pd
from io import BytesIO
from calculation_utils import (calculate_invoicing_table, calculate_by_milestone, 
                               calculate_by_activity, calculate_summary_totals,
                               get_calculation_months)
from activity_utils import MONTHS
from timeline_utils import get_project_config

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

def create_invoicing_excel_report():
    """
    Create an Excel report with invoicing calculations.
    Includes full project duration plus O&M period.
    
    Returns: BytesIO object containing the Excel file
    """
    if not OPENPYXL_AVAILABLE:
        raise ImportError("openpyxl is required for Excel export. Install with: pip install openpyxl")
    
    # Get calculation months (includes O&M period)
    calc_months = get_calculation_months()
    
    # Get calculation data
    invoicing_df = calculate_invoicing_table()
    milestone_df = calculate_by_milestone()
    activity_df = calculate_by_activity()
    summary_totals = calculate_summary_totals()
    
    # Create Excel workbook
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet
    
    # Define styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    
    milestone_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    milestone_font = Font(bold=True, size=10)
    
    total_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    total_font = Font(bold=True, color="000000", size=10)
    
    center_alignment = Alignment(horizontal="center", vertical="center")
    currency_alignment = Alignment(horizontal="right", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Get column letters for dynamic widths
    from openpyxl.utils import get_column_letter
    
    # Sheet 1: Detailed Sub-Milestones
    if not invoicing_df.empty:
        ws1 = wb.create_sheet("Sub-Milestone Details")
        
        # Headers
        for col_idx, col_name in enumerate(['Sub-Milestone'] + calc_months, 1):
            cell = ws1.cell(row=1, column=col_idx, value=col_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_alignment
            cell.border = thin_border
        
        # Data
        for row_idx, (_, row) in enumerate(invoicing_df.iterrows(), 2):
            for col_idx, col_name in enumerate(['Sub-Milestone'] + calc_months, 1):
                value = row.get(col_name, '')
                cell = ws1.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                if col_idx == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = currency_alignment
                    cell.number_format = '#,##0.00'
        
        # Totals row
        totals_row = len(invoicing_df) + 2
        cell = ws1.cell(row=totals_row, column=1, value="TOTAL")
        cell.fill = total_fill
        cell.font = total_font
        cell.border = thin_border
        
        for col_idx, month in enumerate(calc_months, 2):
            cell = ws1.cell(row=totals_row, column=col_idx, value=summary_totals.get(month, 0))
            cell.fill = total_fill
            cell.font = total_font
            cell.border = thin_border
            cell.number_format = '#,##0.00'
            cell.alignment = currency_alignment
        
        # Adjust column widths
        ws1.column_dimensions['A'].width = 25
        for col_idx in range(2, len(calc_months) + 2):
            col_letter = get_column_letter(col_idx)
            ws1.column_dimensions[col_letter].width = 12
    
    # Sheet 2: By Milestone Summary
    if not milestone_df.empty:
        ws2 = wb.create_sheet("By Milestone")
        
        # Headers
        for col_idx, col_name in enumerate(['Milestone'] + calc_months, 1):
            cell = ws2.cell(row=1, column=col_idx, value=col_name)
            cell.fill = milestone_fill
            cell.font = milestone_font
            cell.alignment = center_alignment
            cell.border = thin_border
        
        # Data
        for row_idx, (_, row) in enumerate(milestone_df.iterrows(), 2):
            for col_idx, col_name in enumerate(['Milestone'] + calc_months, 1):
                value = row.get(col_name, '')
                cell = ws2.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                if col_idx == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = currency_alignment
                    cell.number_format = '#,##0.00'
        
        # Adjust column widths
        ws2.column_dimensions['A'].width = 25
        for col_idx in range(2, len(calc_months) + 2):
            col_letter = get_column_letter(col_idx)
            ws2.column_dimensions[col_letter].width = 12
    
    # Sheet 3: By Activity Summary
    if not activity_df.empty:
        ws3 = wb.create_sheet("By Activity")
        
        # Headers
        for col_idx, col_name in enumerate(['Activity'] + calc_months, 1):
            cell = ws3.cell(row=1, column=col_idx, value=col_name)
            cell.fill = milestone_fill
            cell.font = milestone_font
            cell.alignment = center_alignment
            cell.border = thin_border
        
        # Data
        for row_idx, (_, row) in enumerate(activity_df.iterrows(), 2):
            for col_idx, col_name in enumerate(['Activity'] + calc_months, 1):
                value = row.get(col_name, '')
                cell = ws3.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                if col_idx == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = currency_alignment
                    cell.number_format = '#,##0.00'
        
        # Adjust column widths
        ws3.column_dimensions['A'].width = 25
        for col_idx in range(2, len(calc_months) + 2):
            col_letter = get_column_letter(col_idx)
            ws3.column_dimensions[col_letter].width = 12
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_dataframe_to_excel(dataframes_dict):
    """
    Export multiple dataframes to Excel with multiple sheets.
    
    Args:
        dataframes_dict: Dictionary with sheet names as keys and DataFrames as values
    
    Returns: BytesIO object containing the Excel file
    """
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        for sheet_name, df in dataframes_dict.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)
    output.seek(0)
    return output
